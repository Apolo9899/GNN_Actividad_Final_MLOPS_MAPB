"""
Genera la memoria del proyecto MLOps en formato Word (.docx).

Uso:
    pip install python-docx
    python src/generate_report_docx.py
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, color_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def add_code_block(doc, code: str):
    p = doc.add_paragraph()
    p.style = doc.styles["No Spacing"]
    run = p.add_run(code)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    # Fondo gris claro simulado con borde
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        bdr = OxmlElement(f"w:{side}")
        bdr.set(qn("w:val"), "single")
        bdr.set(qn("w:sz"), "4")
        bdr.set(qn("w:space"), "4")
        bdr.set(qn("w:color"), "AAAAAA")
        pBdr.append(bdr)
    pPr.append(pBdr)
    return p


def build_document():
    doc = Document()

    # ── Márgenes ──────────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # ══════════════════════════════════════════════════════════════════════════
    # PORTADA
    # ══════════════════════════════════════════════════════════════════════════
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("GNN Lipophilicity\nProyecto MLOps")
    run.font.size = Pt(24)
    run.font.bold = True

    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Máster Deep Learning — Universidad Politécnica de Madrid\n")
    sub_run = subtitle.add_run("Autor: Marco Apolo Pulpillo Berrocal")
    sub_run.font.bold = True

    doc.add_paragraph()
    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════════
    # 1. INTRODUCCIÓN
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "1. Introducción", level=1)
    add_paragraph(doc, (
        "El presente proyecto aplica metodologías y herramientas de MLOps a un proyecto "
        "de Deep Learning previamente desarrollado durante el Máster. El modelo base es una "
        "Red Neuronal de Grafos (GNN) que predice la lipofilicidad (log P) de moléculas "
        "representadas como grafos atómicos, desarrollado originalmente para el Torneo GNN 2025-26."
    ))
    add_paragraph(doc, (
        "La lipofilicidad es una propiedad fisicoquímica clave en el diseño de fármacos: "
        "determina cómo una molécula se distribuye entre fases acuosas y lipídicas, "
        "afectando directamente a su absorción, distribución y biodisponibilidad."
    ))
    add_paragraph(doc, (
        "El objetivo del proyecto MLOps es transformar el código exploratorio del notebook "
        "original en un sistema de producción fiable, reproducible y mantenible, aplicando "
        "los pilares fundamentales de MLOps: versionado, contenedorización, API REST, "
        "CI/CD automatizado y tracking de experimentos con Weights & Biases."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 2. DATASET
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "2. Dataset", level=1)
    add_paragraph(doc, (
        "El dataset consiste en 4 200 moléculas representadas como grafos en formato GML, "
        "divididas en 3 360 para entrenamiento y 840 para test. Cada grafo representa una "
        "molécula donde:"
    ))
    for item in [
        "Los nodos corresponden a átomos, con 9 features por nodo.",
        "Las aristas representan enlaces químicos entre átomos.",
        "La etiqueta y es el valor de lipofilicidad (variable continua, media ≈ 2.19).",
    ]:
        p = doc.add_paragraph(item, style="List Bullet")

    add_paragraph(doc, (
        "La métrica de evaluación principal es el Mean Squared Error (MSE) sobre el conjunto "
        "de validación (15% del entrenamiento)."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 3. ARQUITECTURA DEL PROYECTO
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "3. Estructura del proyecto", level=1)
    add_paragraph(doc, (
        "Se adoptó la estructura profesional recomendada en clase, separando el código "
        "exploratorio (notebooks) del código de producción (src), los tests y los artefactos:"
    ))
    add_code_block(doc, (
        "gnn-lipophilicity-mlops/\n"
        "├── data/nx_graphs/          # 4 200 grafos GML\n"
        "├── models/                  # Checkpoints entrenados\n"
        "├── notebooks/               # Exploración y experimentación\n"
        "├── src/\n"
        "│   ├── models.py            # Arquitecturas GNN\n"
        "│   ├── dataset.py           # MoleculeDataset (PyG)\n"
        "│   ├── utils.py             # Entrenamiento y evaluación\n"
        "│   ├── train.py             # Script CLI + W&B\n"
        "│   └── inference_api.py     # API FastAPI\n"
        "├── tests/\n"
        "│   ├── test_model.py        # Tests unitarios\n"
        "│   └── test_api.py          # Tests de la API\n"
        "├── .github/workflows/ci.yml # Pipeline CI/CD\n"
        "├── Dockerfile\n"
        "├── requirements.txt\n"
        "└── README.md"
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 4. MODELOS GNN
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "4. Modelos GNN implementados", level=1)
    add_paragraph(doc, (
        "Se implementaron y compararon 5 arquitecturas GNN de complejidad creciente, "
        "todas definidas en src/models.py:"
    ))

    architectures = [
        ("GCNGraph (Exp1)", "Arquitectura baseline. GCN de 3 capas con global_add_pool "
         "y MLP clasificador. Pérdida Huber, sin early stopping."),
        ("GINGraph (Exp2)", "Graph Isomorphism Network con BatchNorm y ReduceLROnPlateau. "
         "Teóricamente más expresivo que GCN gracias a la agregación suma."),
        ("GATGraph (Exp3)", "Graph Attention Network multi-cabeza (4 heads). Aprende a "
         "ponderar la importancia de cada vecino."),
        ("AdvancedGINGraph / GIN-JK (Exp4)", "GIN con conexiones residuales, Jumping Knowledge "
         "(concatena representaciones de todas las capas) y multi-scale pooling "
         "(add + mean + max). Mejor modelo individual."),
        ("VirtualNodeGIN (Exp6)", "GIN-JK con un nodo virtual conectado a todos los nodos "
         "del grafo, que actúa como canal de comunicación global."),
    ]
    for name, desc in architectures:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(name + ": ").bold = True
        p.add_run(desc)

    add_paragraph(doc, (
        "Adicionalmente se implementó un experimento de preentrenamiento auto-supervisado "
        "estilo GraphMAE (Exp8), donde el encoder se preentrenó sobre los 4 200 grafos "
        "(train + test) mediante una tarea de reconstrucción de nodos enmascarados."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 5. EXPERIMENTOS Y RESULTADOS
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "5. Experimentos y resultados", level=1)
    add_paragraph(doc, (
        "Todos los experimentos se ejecutaron con semilla fija (SEED=42) para garantizar "
        "reproducibilidad. Los resultados se resumen en la siguiente tabla:"
    ))

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, txt in enumerate(["Experimento", "Arquitectura", "Val MSE", "Épocas"]):
        hdr[i].text = txt
        hdr[i].paragraphs[0].runs[0].bold = True
        set_cell_bg(hdr[i], "1F5C99")
        hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    rows_data = [
        ("Exp1", "GCN Baseline", "0.7505", "500"),
        ("Exp2", "GIN + BatchNorm", "0.4162", "283 (ES)"),
        ("Exp3", "GAT", "0.6392", "341 (ES)"),
        ("Exp4 ★", "GIN-JK + Multi-scale", "0.3030", "263 (ES)"),
        ("Exp5", "K-Fold Ensemble (×5)", "0.3055 ± 0.011", "—"),
        ("Exp6", "VirtualNodeGIN", "0.3412", "—"),
        ("Exp8", "GraphMAE Fine-Tune", "0.3377", "—"),
        ("Train.py ★★", "GIN-JK (nuevo run)", "0.2808", "269 (ES)"),
    ]
    for exp, arch, mse, epochs in rows_data:
        row = table.add_row().cells
        row[0].text = exp
        row[1].text = arch
        row[2].text = mse
        row[3].text = epochs
        if "★" in exp:
            for cell in row:
                set_cell_bg(cell, "D9EAD3")

    doc.add_paragraph()
    add_paragraph(doc, (
        "El mejor resultado lo obtiene el GIN-JK entrenado con el script train.py (MSE 0.2808), "
        "mejorando al notebook original (0.3030) gracias al early stopping y el scheduler "
        "CosineAnnealingLR bien configurado. ES = Early Stopping."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 6. TRACKING CON WEIGHTS & BIASES
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "6. Tracking de experimentos con Weights & Biases", level=1)
    add_paragraph(doc, (
        "Se integró Weights & Biases (W&B) en el script de entrenamiento para registrar "
        "automáticamente métricas, hiperparámetros y checkpoints. El tracking incluye:"
    ))
    for item in [
        "train_loss y val_mse por época.",
        "Learning rate a lo largo del entrenamiento.",
        "Resumen final: best_val_mse y best_epoch.",
        "Configuración completa del run (arquitectura, hiperparámetros, semilla, dispositivo).",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    add_paragraph(doc, (
        "Adicionalmente, se desarrolló un script (log_experiments_wandb.py) que registra "
        "los 7 experimentos del notebook original como runs independientes, permitiendo "
        "su comparación en el W&B Report."
    ))
    add_paragraph(doc, (
        "El W&B Report incluye: tabla comparativa de todos los runs, curvas de entrenamiento "
        "de los modelos con datos paso a paso, gráfico de barras de MSE por arquitectura "
        "y análisis de conclusiones."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 7. API REST
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "7. API REST con FastAPI", level=1)
    add_paragraph(doc, (
        "El modelo se expone como servicio mediante una API REST implementada con FastAPI "
        "(src/inference_api.py). La API sigue las buenas prácticas vistas en clase:"
    ))
    for item in [
        "Validación de entradas con Pydantic v2 (tipos, rangos, campos obligatorios).",
        "Carga del modelo en el arranque del servidor mediante lifespan context manager.",
        "Manejo estructurado de errores con HTTPException y logging.",
        "Documentación automática en /docs (Swagger UI) y /redoc.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "Endpoints disponibles", level=2)
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = "Table Grid"
    hdr2 = table2.rows[0].cells
    for i, txt in enumerate(["Método", "Ruta", "Descripción"]):
        hdr2[i].text = txt
        hdr2[i].paragraphs[0].runs[0].bold = True
        set_cell_bg(hdr2[i], "1F5C99")
        hdr2[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for method, route, desc in [
        ("GET", "/health", "Estado del servicio y del modelo"),
        ("GET", "/model/info", "Metadatos del modelo cargado"),
        ("POST", "/predict", "Predicción de lipofilicidad"),
        ("GET", "/docs", "Documentación Swagger interactiva"),
    ]:
        row = table2.add_row().cells
        row[0].text = method
        row[1].text = route
        row[2].text = desc

    doc.add_paragraph()
    add_heading(doc, "Formato de entrada (/predict)", level=2)
    add_code_block(doc, (
        'POST /predict\n'
        '{\n'
        '  "nodes": [[f1, f2, ..., f9], ...],  // lista de vectores de 9 features\n'
        '  "edges": [[src, dst], ...]           // lista de aristas\n'
        '}\n\n'
        '// Respuesta:\n'
        '{"lipophilicity": 2.14, "inference_time_ms": 3.5}'
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 8. TESTS
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "8. Tests automatizados", level=1)
    add_paragraph(doc, (
        "Se implementaron 24 tests con pytest, divididos en dos módulos:"
    ))

    add_heading(doc, "test_model.py (17 tests)", level=2)
    for item in [
        "test_output_shape: verifica shape (batch_size, 1) para las 5 arquitecturas.",
        "test_output_dtype: verifica que la salida es float32.",
        "test_gradients_not_none: comprueba que la cabeza MLP recibe gradientes.",
        "test_advanced_gin_overfit_small_batch: el modelo memoriza un lote pequeño (sanity check).",
        "test_single_graph_inference: inferencia con batch_size=1.",
        "test_no_nan_in_output: la salida no contiene NaN.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "test_api.py (7 tests)", level=2)
    for item in [
        "test_health_endpoint: /health responde 200.",
        "test_model_info_endpoint: /model/info devuelve metadatos correctos.",
        "test_predict_valid_input: predicción válida devuelve lipophilicity float.",
        "test_predict_wrong_feature_count: 422 si los nodos no tienen 9 features.",
        "test_predict_invalid_edge: 422 si una arista referencia un nodo inexistente.",
        "test_predict_empty_nodes: 422 con lista de nodos vacía.",
        "test_predict_no_edges: 200 con grafo sin aristas (grafo desconectado).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_paragraph(doc, "Todos los tests se ejecutan con: pytest tests/ -v")

    # ══════════════════════════════════════════════════════════════════════════
    # 9. DOCKER
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "9. Contenedorización con Docker", level=1)
    add_paragraph(doc, (
        "El servicio de inferencia se empaqueta en una imagen Docker basada en "
        "python:3.10-slim con PyTorch en modo CPU (más ligero para inferencia):"
    ))
    add_code_block(doc, (
        "# Construir la imagen\n"
        "docker build -t gnn-lipophilicity:latest .\n\n"
        "# Ejecutar el contenedor\n"
        "docker run -p 8000:8000 gnn-lipophilicity:latest\n\n"
        "# La API queda disponible en http://localhost:8000"
    ))
    add_paragraph(doc, (
        "El Dockerfile incorpora un HEALTHCHECK que verifica el endpoint /health "
        "cada 30 segundos, garantizando que el orquestador detecte fallos del servicio."
    ))

    # ══════════════════════════════════════════════════════════════════════════
    # 10. CI/CD
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "10. Pipeline CI/CD con GitHub Actions", level=1)
    add_paragraph(doc, (
        "El archivo .github/workflows/ci.yml define un pipeline que se ejecuta "
        "automáticamente en cada push a main o develop y en cada Pull Request:"
    ))
    table3 = doc.add_table(rows=1, cols=2)
    table3.style = "Table Grid"
    hdr3 = table3.rows[0].cells
    for i, txt in enumerate(["Job", "Qué hace"]):
        hdr3[i].text = txt
        hdr3[i].paragraphs[0].runs[0].bold = True
        set_cell_bg(hdr3[i], "1F5C99")
        hdr3[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for job, desc in [
        ("test", "Instala dependencias y ejecuta pytest tests/"),
        ("lint", "Verifica estilo de código con flake8 (PEP8)"),
        ("docker", "Construye la imagen Docker (depende de que pasen los tests)"),
    ]:
        row = table3.add_row().cells
        row[0].text = job
        row[1].text = desc

    # ══════════════════════════════════════════════════════════════════════════
    # 11. CONCLUSIONES
    # ══════════════════════════════════════════════════════════════════════════
    add_heading(doc, "11. Conclusiones", level=1)
    add_paragraph(doc, (
        "El proyecto demuestra cómo aplicar el ciclo completo de MLOps a un modelo "
        "de Deep Learning real:"
    ))
    for item in [
        "La modularización del código (src/) facilita el testing, el mantenimiento y el despliegue.",
        "El tracking con W&B permite comparar experimentos de forma objetiva y reproducible.",
        "La contenedorización con Docker garantiza que el servicio funciona igual en cualquier entorno.",
        "El pipeline CI/CD asegura que cada cambio en el código pasa los tests antes de desplegarse.",
        "El mejor modelo (GIN-JK) alcanza MSE 0.2808, mejorando el resultado del notebook original (0.3030) "
        "gracias a un entrenamiento más controlado con early stopping y CosineAnnealingLR.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # ── Guardar ───────────────────────────────────────────────────────────────
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_path = os.path.join(root, "Memoria_Proyecto_MLOps_PulpilloMarco.docx")
    doc.save(out_path)
    print(f"Documento guardado en: {out_path}")


if __name__ == "__main__":
    build_document()
